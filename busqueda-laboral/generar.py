# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datos import FECHA, AVISOS, OFICINAS, BUSQUEDAS, CANALES

OUT = os.path.dirname(os.path.abspath(__file__))
TITULO = "Oportunidades para fotógrafa y maquilladora social"
SUBTITULO = "Zona Norte GBA y CABA lindante · Relevamiento al " + FECHA

NOTA = [
 ("Qué es esto",
  "Relevamiento de avisos y canales activos en los que se busca fotógrafa o maquilladora en Zona Norte del Gran Buenos Aires "
  "(Vicente López, San Isidro, San Fernando, Tigre, Escobar, Pilar) y en zonas de Capital Federal lindantes con Zona Norte "
  "(Belgrano, Núñez, Palermo). Cada fila incluye el link a la fuente de donde se sacó la información."),
 ("Cómo leerlo",
  "«Avisos concretos» son búsquedas puntuales identificadas con nombre de aviso y link. «Oficinas de Empleo» son los portales "
  "municipales, que es donde aparecen las búsquedas chicas de comercios de la zona y donde sí hay teléfono y mail directos. "
  "«Búsquedas guardadas» son URLs ya filtradas por puesto y zona para revisar cada semana. «Canales y directorios» son los lugares "
  "donde los avisos publican contacto directo (WhatsApp o mail) y donde conviene publicar el propio perfil."),
 ("Sobre los teléfonos y mails",
  "Los portales de empleo argentinos (Computrabajo, Bumeran, ZonaJobs, Indeed) no publican el teléfono ni el mail del que contrata: "
  "la postulación es siempre por la plataforma, con cuenta gratuita. Por eso, en esas filas el «medio de contacto» es el link del aviso. "
  "Los contactos telefónicos y de mail verificados que sí figuran en esta lista son los de las oficinas de empleo municipales y los de "
  "los canales del rubro belleza y eventos, donde el contacto sí se publica."),
 ("Limitación del relevamiento",
  "Este relevamiento se armó con búsqueda web. El entorno donde se ejecutó tiene bloqueado el acceso directo a los sitios de empleo, "
  "así que no se pudo abrir cada aviso uno por uno para extraer datos adicionales. Todo lo que figura acá está respaldado por la fuente "
  "linkeada, pero conviene abrir cada link para confirmar que el aviso siga vigente antes de postularse."),
 ("Prioridad sugerida",
  "1) El aviso de Palermo/Belgrano que busca maquilladora Y fotógrafo en la misma producción. 2) ALTHI SRL en Martínez. "
  "3) Alta en los portales municipales de Vicente López y San Isidro, que es donde más aparecen búsquedas chicas de la zona. "
  "4) Peluquería al Día, que publica WhatsApp directo en cada aviso. 5) Perfil de Helohim Pitrelli en Bumeran, que publica seguido."),
]

# ---------------------------------------------------------------- EXCEL
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AZUL   = "1F3864"
AZULC  = "D9E2F3"
GRIS   = "F2F2F2"
BORDE  = Border(*[Side(style="thin", color="BFBFBF")]*4)

def hoja(wb, nombre, headers, filas, anchos, link_col, alturas=None):
    ws = wb.create_sheet(nombre)
    ws.append(headers)
    for c in range(1, len(headers)+1):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(bold=True, color="FFFFFF", size=11)
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cel.border = BORDE
    ws.row_dimensions[1].height = 30

    for i, fila in enumerate(filas, start=2):
        ws.append(list(fila))
        for c in range(1, len(headers)+1):
            cel = ws.cell(row=i, column=c)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
            cel.border = BORDE
            if i % 2 == 0:
                cel.fill = PatternFill("solid", fgColor=GRIS)
        lc = ws.cell(row=i, column=link_col)
        url = str(lc.value or "")
        if url.startswith("http"):
            lc.hyperlink = url
            lc.font = Font(color="0563C1", underline="single", size=9)
        ws.row_dimensions[i].height = (alturas or 60)

    for c, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(filas)+1)
    return ws

wb = Workbook()
ws = wb.active
ws.title = "Leer primero"
ws["A1"] = TITULO
ws["A1"].font = Font(bold=True, size=16, color=AZUL)
ws["A2"] = SUBTITULO
ws["A2"].font = Font(size=11, italic=True, color="595959")
ws["A4"] = "Total de contactos y fuentes relevadas: %d" % (len(AVISOS)+len(OFICINAS)+len(BUSQUEDAS)+len(CANALES))
ws["A4"].font = Font(bold=True, size=11)
r = 6
for titulo, texto in NOTA:
    ws.cell(row=r, column=1, value=titulo).font = Font(bold=True, size=11, color=AZUL)
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=AZULC)
    ws.cell(row=r, column=1).alignment = Alignment(vertical="center")
    c = ws.cell(row=r+1, column=1, value=texto)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r+1].height = 58
    r += 3
ws.column_dimensions["A"].width = 130

H1 = ["Categoría", "Perfil", "Quién busca", "Puesto / aviso", "Zona", "Detalle del aviso", "Contacto directo", "Cómo postularse", "Link a la fuente"]
hoja(wb, "1. Avisos concretos", H1, AVISOS, [15, 12, 34, 34, 26, 58, 40, 46, 62], 9, 78)
hoja(wb, "2. Oficinas de Empleo", H1, OFICINAS, [15, 12, 34, 30, 30, 58, 44, 46, 58], 9, 92)

H3 = ["Portal", "Perfil", "Búsqueda ya filtrada", "Link a la fuente"]
hoja(wb, "3. Búsquedas guardadas", H3, BUSQUEDAS, [18, 13, 46, 88], 4, 22)

H4 = ["Categoría", "Perfil", "Dónde", "Zona", "Por qué sirve", "Contacto directo", "Qué hacer", "Link a la fuente"]
hoja(wb, "4. Canales y directorios", H4, CANALES, [22, 12, 40, 26, 62, 34, 48, 62], 8, 86)

xlsx = os.path.join(OUT, "Contactos-fotografia-maquillaje-ZonaNorte.xlsx")
wb.save(xlsx)
print("XLSX ->", xlsx)

# ---------------------------------------------------------------- PDF
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                Spacer, Table, TableStyle, KeepTogether)

pdf = os.path.join(OUT, "Contactos-fotografia-maquillaje-ZonaNorte.pdf")
PW, PH = landscape(A4)
M = 12*mm
ss = getSampleStyleSheet()
NAVY = colors.HexColor("#1F3864")

def st(name, **kw):
    base = dict(fontName="Helvetica", fontSize=7.4, leading=9.2, textColor=colors.HexColor("#222222"))
    base.update(kw)
    return ParagraphStyle(name, **base)

S_CELL = st("cell")
S_BOLD = st("cellb", fontName="Helvetica-Bold")
S_LINK = st("link", fontSize=6.4, leading=8, textColor=colors.HexColor("#0563C1"))
S_HEAD = st("head", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.white)
S_H1   = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=19, leading=23, textColor=NAVY)
S_H2   = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5, leading=16, textColor=NAVY,
                        spaceBefore=8, spaceAfter=5)
S_SUB  = ParagraphStyle("sub", fontName="Helvetica-Oblique", fontSize=9.5, leading=13,
                        textColor=colors.HexColor("#595959"))
S_BODY = ParagraphStyle("body", fontName="Helvetica", fontSize=9, leading=12.6,
                        textColor=colors.HexColor("#333333"))
S_NOTE = ParagraphStyle("note", fontName="Helvetica-Bold", fontSize=9.3, leading=12.6, textColor=NAVY)

def esc(t):
    return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

def link(u):
    return Paragraph('<link href="%s">%s</link>' % (esc(u), esc(u)), S_LINK)

def tabla(headers, rows, widths, linkidx):
    data = [[Paragraph(esc(h), S_HEAD) for h in headers]]
    for r in rows:
        out = []
        for i, v in enumerate(r):
            if i == linkidx:
                out.append(link(v))
            else:
                out.append(Paragraph(esc(v), S_BOLD if i in (2, 3) and len(headers) > 4 else S_CELL))
        data.append(out)
    t = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    for i in range(1, len(data)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F4F6FA")))
    t.setStyle(TableStyle(style))
    return t

story = []
story.append(Paragraph(TITULO, S_H1))
story.append(Paragraph(SUBTITULO, S_SUB))
story.append(Spacer(1, 5*mm))
story.append(Paragraph("Total de contactos y fuentes relevadas: <b>%d</b> &nbsp;·&nbsp; "
                       "Avisos concretos: <b>%d</b> &nbsp;·&nbsp; Oficinas de Empleo: <b>%d</b> &nbsp;·&nbsp; "
                       "Búsquedas guardadas: <b>%d</b> &nbsp;·&nbsp; Canales y directorios: <b>%d</b>"
                       % (len(AVISOS)+len(OFICINAS)+len(BUSQUEDAS)+len(CANALES),
                          len(AVISOS), len(OFICINAS), len(BUSQUEDAS), len(CANALES)), S_BODY))
story.append(Spacer(1, 4*mm))
for titulo, texto in NOTA:
    story.append(Paragraph(titulo, S_NOTE))
    story.append(Paragraph(texto, S_BODY))
    story.append(Spacer(1, 3*mm))

W = PW - 2*M
story.append(Spacer(1, 3*mm))
story.append(Paragraph("1. Avisos concretos", S_H2))
story.append(tabla(["Perfil", "Quién busca", "Puesto / aviso", "Zona", "Detalle", "Contacto", "Cómo postularse", "Link"],
                   [(a[1], a[2], a[3], a[4], a[5], a[6], a[7], a[8]) for a in AVISOS],
                   [W*0.055, W*0.135, W*0.125, W*0.10, W*0.235, W*0.10, W*0.13, W*0.12], 7))

story.append(Paragraph("2. Oficinas de Empleo municipales de Zona Norte", S_H2))
story.append(Paragraph("Acá están los teléfonos y mails directos verificados. Es la vía más efectiva para las búsquedas "
                       "chicas de comercios y estudios de la zona, que no llegan a los portales grandes.", S_BODY))
story.append(Spacer(1, 2*mm))
story.append(tabla(["Perfil", "Organismo", "Servicio", "Zona que cubre", "Detalle", "Contacto directo", "Qué hacer", "Link"],
                   [(o[1], o[2], o[3], o[4], o[5], o[6], o[7], o[8]) for o in OFICINAS],
                   [W*0.05, W*0.135, W*0.10, W*0.115, W*0.215, W*0.15, W*0.115, W*0.12], 7))

story.append(Paragraph("3. Búsquedas guardadas — revisar una vez por semana", S_H2))
story.append(tabla(["Portal", "Perfil", "Búsqueda ya filtrada por puesto y zona", "Link"],
                   [(b[0], b[1], b[2], b[3]) for b in BUSQUEDAS],
                   [W*0.13, W*0.10, W*0.36, W*0.41], 3))

story.append(Paragraph("4. Canales y directorios donde el contacto sí se publica", S_H2))
story.append(tabla(["Categoría", "Perfil", "Dónde", "Zona", "Por qué sirve", "Contacto", "Qué hacer", "Link"],
                   [(c[0], c[1], c[2], c[3], c[4], c[5], c[6], c[7]) for c in CANALES],
                   [W*0.10, W*0.055, W*0.135, W*0.09, W*0.245, W*0.10, W*0.155, W*0.12], 7))

def deco(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(NAVY)
    canvas.rect(0, PH-6*mm, PW, 6*mm, stroke=0, fill=1)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#777777"))
    canvas.drawString(M, 7*mm, "%s · %s" % (TITULO, SUBTITULO))
    canvas.drawRightString(PW-M, 7*mm, "Página %d" % doc.page)
    canvas.restoreState()

doc = BaseDocTemplate(pdf, pagesize=landscape(A4),
                      leftMargin=M, rightMargin=M, topMargin=M, bottomMargin=13*mm,
                      title=TITULO, author="Relevamiento", subject=SUBTITULO)
frame = Frame(M, 13*mm, W, PH-13*mm-M, id="f")
doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=deco)])
doc.build(story)
print("PDF  ->", pdf)
